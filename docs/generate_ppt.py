import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # Define a clean color palette
    COLOR_BG = RGBColor(248, 250, 252)       # Light gray-white
    COLOR_TITLE = RGBColor(15, 23, 42)       # Dark Slate
    COLOR_ACCENT = RGBColor(37, 99, 235)     # Royal Blue
    COLOR_TEXT = RGBColor(71, 85, 105)       # Slate Gray
    COLOR_BORDER = RGBColor(226, 232, 240)   # Light border
    
    # Theme colors for diagram boxes
    COLOR_CLIENT = RGBColor(13, 148, 136)    # Teal
    COLOR_BACKEND = RGBColor(37, 99, 235)   # Blue
    COLOR_MEDIA = RGBColor(234, 88, 12)      # Orange
    COLOR_EDGE = RGBColor(79, 70, 229)       # Indigo
    
    # Helper to apply background color
    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    # Helper to add standard title
    def add_slide_title(slide, text):
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9.0), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
        
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Calibri"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = COLOR_TITLE
        return title_box

    # Helper to add a formatted content block
    def add_content_block(slide, left, top, width, height, title, items):
        # Add background card shape (rectangle)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left, top, width, height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.color.rgb = COLOR_BORDER
        shape.line.width = Pt(1)
        
        # Add textbox inside
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
        
        # Block Title
        if title:
            p_title = tf.paragraphs[0]
            p_title.text = title
            p_title.font.name = "Calibri"
            p_title.font.size = Pt(20)
            p_title.font.bold = True
            p_title.font.color.rgb = COLOR_ACCENT
            p_title.space_after = Pt(10)
        
        # Block Items
        for idx, item in enumerate(items):
            p = tf.add_paragraph() if (idx > 0 or not title) else tf.paragraphs[0]
            if not title and idx == 0:
                p.text = item
            else:
                p.text = f"•  {item}"
            p.font.name = "Calibri"
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT
            p.space_after = Pt(6)

    # Diagram Helpers
    def draw_node(slide, text, left, top, width, height, bg_color, text_color=RGBColor(255, 255, 255), is_oval=False):
        shape_type = MSO_SHAPE.OVAL if is_oval else MSO_SHAPE.ROUNDED_RECTANGLE
        shape = slide.shapes.add_shape(shape_type, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = COLOR_BORDER
        shape.line.width = Pt(1)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(4)
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = text_color
        return shape

    def draw_arrow(slide, left, top, width, height, text=None, text_offset_y=-22, direction="right", color=RGBColor(148, 163, 184)):
        shape_type = MSO_SHAPE.RIGHT_ARROW if direction == "right" else (MSO_SHAPE.LEFT_ARROW if direction == "left" else MSO_SHAPE.DOWN_ARROW)
        arrow = slide.shapes.add_shape(shape_type, left, top, width, height)
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = color
        arrow.line.fill.background()
        
        if text:
            tb = slide.shapes.add_textbox(left - Inches(0.5), top + Inches(text_offset_y/72.0), width + Inches(1.0), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = text
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Calibri"
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEXT

    # ================= SLIDE 1: Title Slide =================
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(9.0), Inches(3.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Enterprise Video Management System (VMS)"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Calibri"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE
    p.space_after = Pt(14)
    
    p2 = tf.add_paragraph()
    p2.text = "End-to-End Technical Stack & Architecture Overview"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Calibri"
    p2.font.size = Pt(22)
    p2.font.color.rgb = COLOR_ACCENT
    p2.space_after = Pt(20)
    
    p3 = tf.add_paragraph()
    p3.text = "System Integration, Core Responsibilities & Workflows"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = "Calibri"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_TEXT
    
    # ================= SLIDE 2: Core Goals =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "System Goals & Core Objectives")
    
    add_content_block(
        slide, Inches(0.5), Inches(1.5), Inches(4.2), Inches(5.0),
        "Key Capabilities",
        [
            "Low-Latency Live Streaming: Serve sub-second streaming via WebRTC (WHEP) for real-time surveillance monitoring.",
            "Native Quality Recording: Store video natively as MP4 files directly from camera streams to protect file fidelity and disk I/O.",
            "Seamless Fallback: Failover gracefully to Low-Latency HLS (LL-HLS) if WebRTC connection paths are blocked by network firewalls."
        ]
    )
    
    add_content_block(
        slide, Inches(5.3), Inches(1.5), Inches(4.2), Inches(5.0),
        "Resource Optimizations",
        [
            "On-Demand Transcoding: Only transcode H.265 feeds to browser-friendly H.264 when active users are watching.",
            "Shared Transcoder Streams: Spawns one FFmpeg process per camera that is shared across all concurrent viewers.",
            "Viewer-Based Lifecycle: Automatically shuts down transcoders after a 60-second grace period once the last viewer leaves."
        ]
    )

    # ================= SLIDE 3: Architecture Layers =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "VMS High-Level Architecture")
    
    add_content_block(
        slide, Inches(0.5), Inches(1.5), Inches(2.7), Inches(5.0),
        "1. Client Layer",
        [
            "React Web App: Admin tools, config panels, and streaming walls.",
            "WebRTC WHEP Client: Receives raw media streams under 1s latency.",
            "hls.js Fallback: Handles HLS segment parsing for legacy devices."
        ]
    )
    add_content_block(
        slide, Inches(3.65), Inches(1.5), Inches(2.7), Inches(5.0),
        "2. Control Layer",
        [
            "FastAPI App: Exposes APIs and manages WebSocket health feeds.",
            "Transcoder Manager: Oversees external FFmpeg process lifecycles.",
            "TCP Edge Receiver: Listens on Port 9999 for custom ingress feeds."
        ]
    )
    add_content_block(
        slide, Inches(6.8), Inches(1.5), Inches(2.7), Inches(5.0),
        "3. Media & Caching",
        [
            "MediaMTX: Lightweight gateway for RTSP ingress and WebRTC egress.",
            "Redis Session Cache: Synchronizes active transcoders and viewer locks.",
            "Relational Database: Houses schemas for streams and recordings."
        ]
    )

    # ================= SLIDE 4: Frontend Stack =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Frontend Technology Stack")
    
    add_content_block(
        slide, Inches(0.5), Inches(1.5), Inches(4.2), Inches(5.0),
        "Framework & Build System",
        [
            "React (v18.3.1): High-performance UI rendering through components.",
            "TypeScript (v5.6.3): Strong typing that aligns frontend requests with backend schemas.",
            "Vite (v5.4.8): Rapid build tooling using native ES modules.",
            "Vanilla CSS (styles.css): Complete style sheets for dark mode layouts and responsive grids."
        ]
    )
    add_content_block(
        slide, Inches(5.3), Inches(1.5), Inches(4.2), Inches(5.0),
        "Video Playback & Assets",
        [
            "WebRTC (WHEP): Real-time playback integration utilizing native browser APIs.",
            "hls.js (v1.5.15): Media Source Extensions (MSE) player to parse fallbacks.",
            "Lucide React: Modular SVG icons for system dashboards.",
            "Local Storage Cache: Remembers selected streams across page reloads."
        ]
    )

    # ================= SLIDE 5: Backend Stack =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Backend Control Stack")
    
    add_content_block(
        slide, Inches(0.5), Inches(1.5), Inches(4.2), Inches(5.0),
        "FastAPI & Process Core",
        [
            "FastAPI (v0.115.0): Asynchronous Python API framework with routing.",
            "Uvicorn (v0.30.6): ASGI server driving server execution.",
            "Pydantic Settings: Type-safe settings validation from system .env files.",
            "Asyncio Subprocess: Safely spawns background FFmpeg transcoders without blocking CPU execution loops."
        ]
    )
    add_content_block(
        slide, Inches(5.3), Inches(1.5), Inches(4.2), Inches(5.0),
        "Asynchronous Clients",
        [
            "HTTPX: Asynchronous requests targeting MediaMTX control APIs and camera registries.",
            "aiofiles: Prevents disk blocking during timeline generation and clip export operations.",
            "Custom TCP Server: Standard sockets parsing camera framing payloads on Port 9999."
        ]
    )

    # ================= SLIDE 6: Media Infrastructure =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Media & Streaming Infrastructure")
    
    add_content_block(
        slide, Inches(0.5), Inches(1.5), Inches(4.2), Inches(5.0),
        "Media Server & Transcoding",
        [
            "MediaMTX (v1.9.0): Zero-dependency media processor. Handles RTSP streams, segments storage, and WebRTC/HLS publishers.",
            "FFmpeg Transcoders: Convert H.265 video packets into H.264 streams. Uses speed parameters like 'ultrafast' for real-time delivery.",
            "Hardware Acceleration: Configurable to offload CPU load using NVIDIA NVENC or Intel QSV hardware codecs."
        ]
    )
    add_content_block(
        slide, Inches(5.3), Inches(1.5), Inches(4.2), Inches(5.0),
        "NAT Traversal & TCP Ingress",
        [
            "Coturn (TURN/STUN): Resolves public ICE candidates. Relays media packages for players operating across separate subnets.",
            "Custom TCP Receiver: Built with asyncio to read custom framing headers. Extracts VPS, SPS, PPS, and NAL units, then remuxes raw streams into MediaMTX via RTSP push."
        ]
    )

    # ================= SLIDE 7: Database & Caching =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Storage, Databases & Caching")
    
    add_content_block(
        slide, Inches(0.5), Inches(1.5), Inches(4.2), Inches(5.0),
        "Database Engine (ORM)",
        [
            "SQLAlchemy (v2.0.34): Coordinates relational schema mappings and async data queries.",
            "Alembic (v1.13.1): Tracks schema migrations across versions.",
            "SQLite & aiosqlite: Default database for local dev and small deployments.",
            "PostgreSQL & asyncpg: Production relational engine supporting async connection pooling."
        ]
    )
    add_content_block(
        slide, Inches(5.3), Inches(1.5), Inches(4.2), Inches(5.0),
        "State Caching & Session Sync",
        [
            "Redis (v5.0.3): Fast key-value store for cross-process synchronization and distributed lock management.",
            "Redis Viewer Tracker: Monitors active user connections to trigger real-time H.265 transcoding start and stop instructions.",
            "Upstream Caching Fallback: Caches remote camera configs into local JSON formats if settings APIs timeout."
        ]
    )

    # ================= SLIDE 8: Edge-Push Agent =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Edge Push Agent (edge-push)")
    
    add_content_block(
        slide, Inches(0.5), Inches(1.5), Inches(4.2), Inches(5.0),
        "Core Capabilities",
        [
            "Standalone Client: Runs on local gate hardware to push camera video to remote servers.",
            "Automatic Detection: Uses ffprobe to identify camera codec profiles and frame rates.",
            "Zero-Reencoding Push: Copies RTSP streams directly, minimizing CPU overhead on edge devices."
        ]
    )
    add_content_block(
        slide, Inches(5.3), Inches(1.5), Inches(4.2), Inches(5.0),
        "Downsampling & Protocol",
        [
            "Downsampling (compressedFps): Filters and drops frames at the edge to reduce network load.",
            "Custom TCP protocol: Bundles configuration parameters (VPS/SPS/PPS) and NAL units into binary packets.",
            "Network Fault Tolerance: Keeps trying to reconnect when connections drop."
        ]
    )

    # ================= SLIDE 9: Production DevOps =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Production DevOps & Deployment")
    
    add_content_block(
        slide, Inches(0.5), Inches(1.5), Inches(4.2), Inches(5.0),
        "Production Operations",
        [
            "Docker Compose: Bundles the entire VMS stack (FastAPI, MediaMTX, Redis, DB) for simple deployment.",
            "Nginx Web Server: Acts as the primary HTTP reverse proxy. Serves built static frontend files.",
            "WebSocket Upgrades: Configured to handle WebSocket handshakes and state synchronization."
        ]
    )
    add_content_block(
        slide, Inches(5.3), Inches(1.5), Inches(4.2), Inches(5.0),
        "Host Configuration",
        [
            "systemd Services: Registers VMS units as background services with automatic restart behaviors.",
            "UFW Firewall: Secures unnecessary ports while keeping ports 80/443 (UI), 8000 (API), 9999 (TCP Push), 8554 (RTSP), 8889 (WebRTC), and 3478 (TURN) open."
        ]
    )

    # ================= SLIDE 10: DIAGRAM - WebRTC Signaling Flow =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Flow Diagram: WebRTC WHEP Signaling")
    
    # Draw boxes
    draw_node(slide, "React Frontend\n(WHEP Player)", Inches(0.5), Inches(2.2), Inches(2.2), Inches(1.2), COLOR_CLIENT)
    draw_node(slide, "FastAPI Backend\n(Signaling Proxy)", Inches(3.9), Inches(2.2), Inches(2.2), Inches(1.2), COLOR_BACKEND)
    draw_node(slide, "MediaMTX Server\n(WebRTC Core)", Inches(7.3), Inches(2.2), Inches(2.2), Inches(1.2), COLOR_MEDIA)
    
    # Draw arrows
    # Offer: Client -> FastAPI
    draw_arrow(slide, Inches(2.8), Inches(2.5), Inches(1.0), Inches(0.15), "1. WHEP Offer", text_offset_y=-18)
    # Offer: FastAPI -> MediaMTX
    draw_arrow(slide, Inches(6.2), Inches(2.5), Inches(1.0), Inches(0.15), "2. Proxy Offer", text_offset_y=-18)
    # Answer: MediaMTX -> FastAPI
    draw_arrow(slide, Inches(6.2), Inches(3.0), Inches(1.0), Inches(0.15), "3. WHEP Answer", text_offset_y=16, direction="left")
    # Answer: FastAPI -> Client
    draw_arrow(slide, Inches(2.8), Inches(3.0), Inches(1.0), Inches(0.15), "4. Proxy Answer", text_offset_y=16, direction="left")
    
    # Direct Media flow (bottom curved/wide connector)
    media_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(4.5), Inches(9.0), Inches(1.0))
    media_box.fill.solid()
    media_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    media_box.line.color.rgb = COLOR_MEDIA
    media_box.line.width = Pt(2)
    tf = media_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "5. Direct WebRTC Media Stream (UDP / RTP)\nEstablishes peer-to-peer playback connection. Fallbacks use Coturn TURN/STUN relay on Port 3478."
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Calibri"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TITLE

    # ================= SLIDE 11: DIAGRAM - Transcoding Flow =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Flow Diagram: H.265 On-Demand Transcoding")
    
    # Draw sequence horizontal flow
    draw_node(slide, "1. Play Request\n(Browser H.265)", Inches(0.4), Inches(2.5), Inches(1.8), Inches(1.0), COLOR_CLIENT)
    draw_arrow(slide, Inches(2.25), Inches(2.9), Inches(0.4), Inches(0.12))
    
    draw_node(slide, "2. Check Sessions\n(FastAPI / Redis)", Inches(2.7), Inches(2.5), Inches(1.8), Inches(1.0), COLOR_BACKEND)
    draw_arrow(slide, Inches(4.55), Inches(2.9), Inches(0.4), Inches(0.12))
    
    draw_node(slide, "3. Spawn Worker\n(FFmpeg Process)", Inches(5.0), Inches(2.5), Inches(1.8), Inches(1.0), COLOR_MEDIA)
    draw_arrow(slide, Inches(6.85), Inches(2.9), Inches(0.4), Inches(0.12))
    
    draw_node(slide, "4. Remux Egress\n(H.264 WebRTC)", Inches(7.3), Inches(2.5), Inches(1.8), Inches(1.0), COLOR_MEDIA)
    
    # Transcoder details block below
    details_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(4.2), Inches(9.2), Inches(1.8))
    details_box.fill.solid()
    details_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    details_box.line.color.rgb = COLOR_BORDER
    details_box.line.width = Pt(1)
    tf = details_box.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = Inches(0.2)
    
    p1 = tf.paragraphs[0]
    p1.text = "Transcoding Lifecycle & Optimization Details:"
    p1.font.bold = True
    p1.font.size = Pt(14)
    p1.font.color.rgb = COLOR_ACCENT
    p1.space_after = Pt(6)
    
    p2 = tf.add_paragraph()
    p2.text = "• Multiplexing: If additional viewers request the same camera, the system hooks them to the existing running FFmpeg process.\n• CPU Safeguard: Concurrent transcoder count capped at 10 (configurable in env via MAX_ACTIVE_TRANSCODERS).\n• Grace Period: When viewer count hits 0, a 60-second timer begins. If no connections occur, the FFmpeg process is killed."
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLOR_TEXT
    p2.space_after = Pt(4)

    # ================= SLIDE 12: DIAGRAM - Edge Push Ingress Flow =================
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    add_slide_title(slide, "Flow Diagram: TCP Edge Stream Ingress")
    
    # Horizontal flow nodes
    draw_node(slide, "Local IP Camera\n(RTSP Feed)", Inches(0.4), Inches(2.2), Inches(1.8), Inches(1.0), COLOR_TEXT)
    draw_arrow(slide, Inches(2.25), Inches(2.6), Inches(0.4), Inches(0.12))
    
    draw_node(slide, "Edge-Push Agent\n(TCP Packetizer)", Inches(2.7), Inches(2.2), Inches(1.8), Inches(1.0), COLOR_EDGE)
    draw_arrow(slide, Inches(4.55), Inches(2.6), Inches(0.4), Inches(0.12), "Port 9999")
    
    draw_node(slide, "TCP Receiver\n(FastAPI Server)", Inches(5.0), Inches(2.2), Inches(1.8), Inches(1.0), COLOR_BACKEND)
    draw_arrow(slide, Inches(6.85), Inches(2.6), Inches(0.4), Inches(0.12))
    
    draw_node(slide, "MediaMTX\n(RTSP Ingest)", Inches(7.3), Inches(2.2), Inches(1.8), Inches(1.0), COLOR_MEDIA)
    
    # Description block
    desc_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(3.8), Inches(9.2), Inches(2.4))
    desc_box.fill.solid()
    desc_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    desc_box.line.color.rgb = COLOR_BORDER
    desc_box.line.width = Pt(1)
    tf = desc_box.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = Inches(0.2)
    
    p = tf.paragraphs[0]
    p.text = "Ingress Pipeline Breakdown:"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(6)
    
    items = [
        "1. Camera Capture: Edge client hooks into camera's local RTSP stream using subprocess ffmpeg.",
        "2. Binary Packetizer: Raw video NAL units (H.264/H.265) are parsed. Stream headers (VPS/SPS/PPS) are extracted and base64-encoded.",
        "3. TCP socket Push: Connects to central server IP on port 9999 and transmits serialized custom frames.",
        "4. Server Receiver: The backend tcp_receiver loops incoming packets, validates checksums, and launches FFmpeg to remux raw NAL units into MediaMTX via RTSP push.",
        "5. Playback Ready: Once published to MediaMTX, browsers can stream it immediately via WebRTC or LL-HLS."
    ]
    for item in items:
        p_item = tf.add_paragraph()
        p_item.text = item
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = COLOR_TEXT
        p_item.space_after = Pt(3)

    # Save presentation
    output_path = "c:\\Users\\santhosh\\Downloads\\vs\\camera_video_platform_fullstack\\docs\\vms_tech_stack_diagrams.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()
